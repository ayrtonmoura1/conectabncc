# -*- coding: utf-8 -*-
"""
construir_relacoes_finais.py
Gera a versão definitiva e rigorosa das relações entre habilidades da BNCC:
- Progressão Cognitiva (Taxonomia de Bloom: verbos cognitivos e complexidade)
- Lógica Temporal e de Séries (Requisitos de anos anteriores/início; Prepara para anos posteriores/fim)
- Articulação Interdisciplinar e Transversal (outros componentes no mesmo ano/ciclo)
- Extração de Conhecimentos Prévios Estruturantes da Progressão Curricular da BNCC
- Regra Estrita Anti-Alucinação (somente códigos válidos do CSV original de 1.517 habilidades)
- Justificativa pedagógica concisa (Chain of Thought) para cada relação

Saída: relacoes_bncc_final.json
"""

import json
import os
import re
import openpyxl
import pandas as pd
import numpy as np

CSV_PATH = os.path.join('recursos_adicionais', 'Cópia de habilidades bncc - bncc.csv')
EXCEL_PROGRESSAO = os.path.join('recursos_adicionais', 'MapasDeFocoBncc_Unificados.xlsx')
VETORES_JSON = 'bncc_vetores.json'
OUTPUT_JSON = 'relacoes_bncc_final.json'

# ======================================================================
# 1. Taxonomia de Bloom (Classificação de Verbos Cognitivos)
# ======================================================================
BLOOM_TAXONOMY = {
    1: {
        'nome': 'Lembrar / Identificar',
        'verbos': {
            'identificar', 'reconhecer', 'nomear', 'listar', 'apontar', 'localizar',
            'recordar', 'reproduzir', 'relatar', 'registrar', 'descrever', 'ler',
            'observar', 'copiar', 'dizer', 'selecionar', 'perceber'
        }
    },
    2: {
        'nome': 'Compreender',
        'verbos': {
            'compreender', 'entender', 'classificar', 'explicar', 'exemplificar',
            'inferir', 'interpretar', 'resumir', 'distinguir', 'caracterizar',
            'associar', 'agrupar', 'ordenar', 'comparar', 'confrontar', 'discutir',
            'escutar', 'apreciar', 'diferenciar'
        }
    },
    3: {
        'nome': 'Aplicar',
        'verbos': {
            'aplicar', 'resolver', 'calcular', 'utilizar', 'usar', 'demonstrar',
            'executar', 'medir', 'estimar', 'operar', 'manusear', 'experimentar',
            'ilustrar', 'traçar', 'escrever', 'recitar', 'cantar', 'compartilhar',
            'participar', 'praticar', 'pesquisar', 'coletar', 'consultar', 'adotar'
        }
    },
    4: {
        'nome': 'Analisar',
        'verbos': {
            'analisar', 'organizar', 'decompor', 'examinar', 'correlacionar',
            'relacionar', 'investigar', 'debater', 'contrastar', 'discriminar',
            'categorizar', 'problematizar', 'sistematizar'
        }
    },
    5: {
        'nome': 'Avaliar / Argumentar',
        'verbos': {
            'avaliar', 'julgar', 'criticar', 'justificar', 'validar', 'argumentar',
            'ponderar', 'posicionar', 'posicionar-se', 'defender', 'concluir',
            'contrapor', 'verificar'
        }
    },
    6: {
        'nome': 'Criar / Elaborar',
        'verbos': {
            'criar', 'produzir', 'construir', 'planejar', 'elaborar', 'desenvolver',
            'projetar', 'propor', 'formular', 'compor', 'retextualizar', 'encenar',
            'inventar', 'transformar', 'editar', 'diagramar', 'revisar'
        }
    }
}

VERBO_PARA_BLOOM = {}
for level, data in BLOOM_TAXONOMY.items():
    for v in data['verbos']:
        VERBO_PARA_BLOOM[v] = level

def extrair_verbo_bloom(texto):
    if not texto:
        return 'identificar', 1
    t = str(texto).strip()
    t = re.sub(r'^\([A-Z0-9]+\)\s*', '', t)
    palavras = re.findall(r'[a-zA-ZÀ-ÿ]+', t.lower())
    for w in palavras[:4]:
        if w in VERBO_PARA_BLOOM:
            return w, VERBO_PARA_BLOOM[w]
        if w.endswith(('ar', 'er', 'ir', 'or')):
            return w, 2
    return palavras[0] if palavras else 'identificar', 1

# ======================================================================
# 2. Normalização de Componente e Temporalidade
# ======================================================================
MAPA_TEMPORAL = {
    'EI01': (0.1, 0.1, 'Bebês (0 a 1a 6m)', 'EI'),
    'EI02': (0.2, 0.2, 'Crianças bem pequenas (1a 7m a 3a 11m)', 'EI'),
    'EI03': (0.3, 0.3, 'Crianças pequenas (4 a 5a 11m)', 'EI'),
    'EF01': (1.0, 1.0, '1º ano', 'EF_INICIAIS'),
    'EF02': (2.0, 2.0, '2º ano', 'EF_INICIAIS'),
    'EF12': (1.0, 2.0, '1º e 2º anos', 'EF_INICIAIS'),
    'EF03': (3.0, 3.0, '3º ano', 'EF_INICIAIS'),
    'EF04': (4.0, 4.0, '4º ano', 'EF_INICIAIS'),
    'EF05': (5.0, 5.0, '5º ano', 'EF_INICIAIS'),
    'EF35': (3.0, 5.0, '3º ao 5º ano', 'EF_INICIAIS'),
    'EF15': (1.0, 5.0, '1º ao 5º ano', 'EF_INICIAIS'),
    'EF06': (6.0, 6.0, '6º ano', 'EF_FINAIS'),
    'EF07': (7.0, 7.0, '7º ano', 'EF_FINAIS'),
    'EF67': (6.0, 7.0, '6º e 7º anos', 'EF_FINAIS'),
    'EF08': (8.0, 8.0, '8º ano', 'EF_FINAIS'),
    'EF09': (9.0, 9.0, '9º ano', 'EF_FINAIS'),
    'EF89': (8.0, 9.0, '8º e 9º anos', 'EF_FINAIS'),
    'EF69': (6.0, 9.0, '6º ao 9º ano', 'EF_FINAIS'),
    'EM13': (10.0, 12.0, '1º ao 3º ano do Ensino Médio', 'EM')
}

def normalizar_componente_e_area(cod, comp_csv):
    cod = str(cod).strip().upper()
    c = str(comp_csv).strip() if pd.notna(comp_csv) and str(comp_csv).lower() != 'nan' else ''
    
    # 1. Educação Infantil
    if cod.startswith('EI'):
        return 'Educação Infantil', 'Educação Infantil'
    
    # 2. Ensino Médio
    if cod.startswith('EM'):
        if 'MAT' in cod:
            return 'Matemática e suas Tecnologias', 'Matemática'
        if 'CNT' in cod:
            return 'Ciências da Natureza e suas Tecnologias', 'Ciências da Natureza'
        if 'CHS' in cod:
            return 'Ciências Humanas e Sociais Aplicadas', 'Ciências Humanas'
        if 'LGG' in cod:
            return 'Linguagens e suas Tecnologias', 'Linguagens'
        return c or 'Ensino Médio', 'Ensino Médio'

    # 3. Ensino Fundamental
    if 'MAT' in cod:
        return 'Matemática', 'Matemática'
    if 'LP' in cod:
        return 'Língua Portuguesa', 'Linguagens'
    if 'CI' in cod:
        return 'Ciências', 'Ciências da Natureza'
    if 'HI' in cod:
        return 'História', 'Ciências Humanas'
    if 'GE' in cod:
        return 'Geografia', 'Ciências Humanas'
    if 'AR' in cod:
        return 'Arte', 'Linguagens'
    if len(cod) >= 6 and cod[4:6] == 'EF':
        return 'Educação Física', 'Linguagens'
    if 'LI' in cod:
        return 'Língua Inglesa', 'Linguagens'

    return c or 'Geral', 'Geral'

def obter_temporalidade(codigo, ano_faixa_str):
    pref = (codigo or '')[:4].upper()
    if pref in MAPA_TEMPORAL:
        return MAPA_TEMPORAL[pref]
    if pref.startswith('EI'):
        return (0.2, 0.2, 'Educação Infantil', 'EI')
    if pref.startswith('EM'):
        return (10.0, 12.0, '1º ao 3º ano do Ensino Médio', 'EM')
    return (5.0, 5.0, ano_faixa_str or 'Ensino Fundamental', 'EF_INICIAIS')

def extrair_numero_sequencia(codigo):
    match = re.search(r'(\d+)$', codigo or '')
    if match:
        try:
            return int(match.group(1))
        except:
            return 99
    return 99

# ======================================================================
# 3. Extração dos Conhecimentos Prévios da Matriz de Progressão BNCC
# ======================================================================
def extrair_matriz_progressao_curricular(caminho_excel, codigos_validos):
    print(f'-> Lendo matrizes de progressão curricular: {caminho_excel}...')
    if not os.path.exists(caminho_excel):
        print(f'   [AVISO] Arquivo {caminho_excel} não encontrado.')
        return {}, {}

    wb = openpyxl.load_workbook(caminho_excel, data_only=True)
    matriz_requisitos = {}
    code_pattern = re.compile(r'\b(E[IFM]\d{2}[A-Z]{2,3}\d{2})\b')

    for sname in wb.sheetnames:
        s = wb[sname]
        col_prev = None
        col_cod = None
        col_obj = None
        header_row = None

        for r in range(1, 6):
            for c in range(1, 15):
                val = str(s.cell(r, c).value or '').lower()
                if 'conhecimento pr' in val or 'habilidades de anos anteriores' in val:
                    col_prev = c
                    header_row = r
                elif 'código da habilidade' in val or 'código' in val and col_cod is None:
                    col_cod = c
                elif 'objetivos de aprendizagem' in val:
                    col_obj = c

            if col_prev and col_cod:
                break

        if not (col_prev and col_cod and header_row):
            continue

        for r in range(header_row + 1, s.max_row + 1):
            val_cod = s.cell(r, col_cod).value
            val_prev = s.cell(r, col_prev).value

            if not val_cod:
                continue

            targets = code_pattern.findall(str(val_cod).upper())
            if not targets:
                continue
            target = targets[0]
            if target not in codigos_validos:
                continue

            if val_prev:
                prev_codes = code_pattern.findall(str(val_prev).upper())
                valid_prev = [pc for pc in prev_codes if pc in codigos_validos and pc != target]
                if valid_prev:
                    if target not in matriz_requisitos:
                        matriz_requisitos[target] = []
                    for pc in valid_prev:
                        if not any(x[0] == pc for x in matriz_requisitos[target]):
                            just = 'Conhecimento prévio estruturante e base conceitual necessária para o desenvolvimento cognitivo da habilidade'
                            matriz_requisitos[target].append((pc, just))

    matriz_prepara_para = {}
    for tgt, lista in matriz_requisitos.items():
        for src, _ in lista:
            if src not in matriz_prepara_para:
                matriz_prepara_para[src] = []
            if not any(x[0] == tgt for x in matriz_prepara_para[src]):
                just = f'Habilidade de suporte formativo que prepara diretamente para o desenvolvimento da habilidade subsequente {tgt}'
                matriz_prepara_para[src].append((tgt, just))

    print(f'   -> Extraídos conhecimentos prévios para {len(matriz_requisitos)} habilidades.')
    print(f'   -> Derivadas progressões para {len(matriz_prepara_para)} habilidades.')
    return matriz_requisitos, matriz_prepara_para

# ======================================================================
# 4. Pipeline Principal de Construção e Avaliação Cognitiva
# ======================================================================
def main():
    print('=' * 80)
    print('CONSTRUÇÃO DAS RELAÇÕES DEFINITIVAS DA BNCC (relacoes_bncc_final.json)')
    print('=' * 80)

    # 1. Carrega CSV Original
    print(f'1. Carregando habilidades do CSV: {CSV_PATH}...')
    df = pd.read_csv(CSV_PATH)
    print(f'   Total de registros brutos: {len(df):,}')

    habilidades = {}
    valid_codes = set()

    for _, row in df.iterrows():
        cod = str(row['COD']).strip().upper()
        valid_codes.add(cod)

        desc = str(row['HABILIDADES']).strip() if pd.notna(row['HABILIDADES']) else ''
        ano_faixa = str(row['ANO/FAIXA']).strip() if pd.notna(row['ANO/FAIXA']) else ''
        ut = str(row['PRÁTICAS DE LINGUAGEM - UNIDADE TEMÁTICA']).strip() if pd.notna(row['PRÁTICAS DE LINGUAGEM - UNIDADE TEMÁTICA']) else ''
        oc_raw = str(row['OBJETOS DE CONHECIMENTO']).strip() if pd.notna(row['OBJETOS DE CONHECIMENTO']) else ''
        coment = str(row['COMENTÁRIO']).strip() if pd.notna(row['COMENTÁRIO']) else ''
        possib = str(row['POSSIBILIDADES PARA O CURRÍCULO']).strip() if pd.notna(row['POSSIBILIDADES PARA O CURRÍCULO']) else ''

        comp, area = normalizar_componente_e_area(cod, row['Disciplina/Area'])

        oc_list = []
        if oc_raw and oc_raw.lower() != 'nan':
            partes = re.split(r'[;\n•\r]+', oc_raw)
            for p in partes:
                limpo = p.strip(' -*•\t')
                if len(limpo) >= 3:
                    oc_list.append(limpo)
        if not oc_list and ut:
            oc_list = [ut]

        t_min, t_max, rotulo_ano, etapa = obter_temporalidade(cod, ano_faixa)
        verbo, bloom_lvl = extrair_verbo_bloom(desc)
        seq_num = extrair_numero_sequencia(cod)

        habilidades[cod] = {
            'codigo': cod,
            'descricao': desc,
            'componente': comp,
            'area': area,
            'ano_fase': rotulo_ano if rotulo_ano else ano_faixa,
            'ano_faixa_orig': ano_faixa,
            'etapa': etapa,
            't_min': t_min,
            't_max': t_max,
            'seq_num': seq_num,
            'verbo': verbo,
            'bloom_lvl': bloom_lvl,
            'unidade_tematica': ut,
            'objetos_conhecimento': oc_list,
            'comentario': coment,
            'possibilidades': possib
        }

    print(f'   Total de habilidades indexadas na memória: {len(habilidades)}')

    # 2. Carrega Vetores pré-computados para Similaridade Semântica Densa
    vetores_map = {}
    if os.path.exists(VETORES_JSON):
        print(f'2. Carregando embeddings densos de {VETORES_JSON}...')
        with open(VETORES_JSON, 'r', encoding='utf-8') as f:
            v_data = json.load(f)
            for item in v_data:
                c = item.get('codigo')
                v = item.get('vetor')
                if c and v:
                    vetores_map[c] = np.array(v, dtype=np.float32)
        print(f'   Embeddings carregados para {len(vetores_map)} habilidades.')
    else:
        print('   [AVISO] bncc_vetores.json não encontrado.')

    # 3. Extrai Matriz de Progressão Curricular Estruturante
    matriz_pre, matriz_pos = extrair_matriz_progressao_curricular(EXCEL_PROGRESSAO, valid_codes)

    # 4. Agrupamento em Listas de Candidatos
    por_componente = {}
    por_etapa = {}
    for cod, h in habilidades.items():
        por_componente.setdefault(h['componente'], []).append(cod)
        por_etapa.setdefault(h['etapa'], []).append(cod)

    def calcular_sim(c1, c2):
        v1 = vetores_map.get(c1)
        v2 = vetores_map.get(c2)
        if v1 is not None and v2 is not None:
            return float(np.dot(v1, v2))
        return 0.0

    def calcular_afinidade_tematica(hA, hB):
        score = 0.0
        detalhes = []

        # Mesma unidade temática
        utA = (hA['unidade_tematica'] or '').lower().strip()
        utB = (hB['unidade_tematica'] or '').lower().strip()
        if utA and utB:
            if utA == utB:
                score += 0.35
                detalhes.append(f'mesma unidade temática ({hA["unidade_tematica"]})')
            elif any(w in utB for w in utA.split() if len(w) > 4):
                score += 0.18
                detalhes.append('unidade temática correlata')

        # Sobreposição em objetos de conhecimento
        ocsA = set(o.lower() for o in hA['objetos_conhecimento'])
        ocsB = set(o.lower() for o in hB['objetos_conhecimento'])
        inter_oc = ocsA.intersection(ocsB)
        if inter_oc:
            score += 0.40
            detalhes.append(f'mesmo objeto de conhecimento ({list(inter_oc)[0]})')
        else:
            palavrasA = set()
            for o in ocsA:
                palavrasA.update(w for w in re.findall(r'[a-zA-ZÀ-ÿ]{4,}', o))
            palavrasB = set()
            for o in ocsB:
                palavrasB.update(w for w in re.findall(r'[a-zA-ZÀ-ÿ]{4,}', o))
            comuns = palavrasA.intersection(palavrasB) - {'para', 'sobre', 'com', 'forma', 'estudo', 'parte'}
            if len(comuns) >= 2:
                score += 0.22
                detalhes.append(f'termos conceituais correlatos ({", ".join(list(comuns)[:2])})')

        # Similaridade vetorial densa
        sim = calcular_sim(hA['codigo'], hB['codigo'])
        score += sim * 0.45

        return score, detalhes, sim

    # ==================================================================
    # 5. Processamento dos 1.517 Registros
    # ==================================================================
    print('3. Processando progressão cognitiva e articulações para todas as habilidades...')
    resultado_final = {}

    total_requisitos = 0
    total_prepara = 0
    total_relacionadas = 0

    count = 0
    for cod_alvo, alvo in habilidades.items():
        count += 1
        if count % 250 == 0 or count == len(habilidades):
            print(f'   -> Processadas {count}/{len(habilidades)} habilidades ({(count/len(habilidades)*100):.1f}%)...')

        comp = alvo['componente']
        t_min = alvo['t_min']
        t_max = alvo['t_max']
        bloom_alvo = alvo['bloom_lvl']
        verbo_alvo = alvo['verbo']

        # --------------------------------------------------------------
        # A) REQUISITOS (Máx 3)
        # --------------------------------------------------------------
        candidatos_requisitos = []
        codigos_candidatos = por_componente.get(comp, [])

        for c_cod in codigos_candidatos:
            if c_cod == cod_alvo:
                continue
            cand = habilidades[c_cod]

            # REGRA TEMPORAL INQUEBRÁVEL:
            # Deve ser de ano ANTERIOR ou INÍCIO do mesmo ano
            if cand['t_min'] > t_min:
                continue  # Nunca pode ser de ano posterior
            if cand['t_max'] > t_max:
                continue  # Não pode extrapolar o ano da alvo

            is_same_year = (cand['t_min'] == t_min and cand['t_max'] == t_max)
            if is_same_year:
                # No mesmo ano, um pré-requisito deve ser estritamente preliminar na sequência curricular
                if cand['seq_num'] >= alvo['seq_num']:
                    continue

            score_afinidade, detalhes, sim = calcular_afinidade_tematica(cand, alvo)

            is_matriz_estruturante = any(r[0] == c_cod for r in matriz_pre.get(cod_alvo, []))
            if is_matriz_estruturante:
                score_afinidade += 2.0
                detalhes.insert(0, 'conhecimento prévio estruturante mapeado na progressão curricular')

            if cand['bloom_lvl'] <= bloom_alvo:
                score_afinidade += 0.15

            if score_afinidade >= 0.50 or is_matriz_estruturante:
                motivo = ''
                if is_matriz_estruturante:
                    motivo = f'Base curricular estruturante: consolida o conceito prévio de "{cand["objetos_conhecimento"][0] if cand["objetos_conhecimento"] else cand["unidade_tematica"]}" para permitir a progressão rumo à habilidade {cod_alvo}.'
                else:
                    dif_anos = t_min - cand['t_min']
                    if dif_anos > 0:
                        motivo = f'Progressão vertical ({cand["ano_fase"]} -> {alvo["ano_fase"]}): constrói a base cognitiva de "{cand["verbo"]}" em {cand["objetos_conhecimento"][0] if cand["objetos_conhecimento"] else "conteúdo correlato"}, necessária para {verbo_alvo} no {alvo["ano_fase"]}.'
                    else:
                        motivo = f'Pré-requisito no mesmo ciclo ({alvo["ano_fase"]}): introduz o fundamento conceitual preliminar de {cand["objetos_conhecimento"][0] if cand["objetos_conhecimento"] else "base temática"} antes do aprofundamento de {cod_alvo}.'

                candidatos_requisitos.append({
                    'codigo': c_cod,
                    'score': score_afinidade,
                    'justificativa': motivo
                })

        candidatos_requisitos.sort(key=lambda x: x['score'], reverse=True)
        final_requisitos = [{'codigo': item['codigo'], 'justificativa': item['justificativa']} for item in candidatos_requisitos[:3]]
        codigos_requisitos_set = set(item['codigo'] for item in final_requisitos)

        # --------------------------------------------------------------
        # B) PREPARA_PARA (Máx 3)
        # --------------------------------------------------------------
        candidatos_prepara = []

        for c_cod in codigos_candidatos:
            if c_cod == cod_alvo or c_cod in codigos_requisitos_set:
                continue
            cand = habilidades[c_cod]

            # REGRA TEMPORAL INQUEBRÁVEL:
            # Deve ser de ano POSTERIOR ou FINAL do mesmo ano
            if cand['t_max'] < t_max:
                continue  # Nunca pode ser de ano anterior
            if cand['t_min'] < t_min:
                continue  # Não pode começar antes do ano da alvo

            is_same_year = (cand['t_min'] == t_min and cand['t_max'] == t_max)
            if is_same_year:
                # No mesmo ano, deve ser estritamente posterior na sequência curricular
                if cand['seq_num'] <= alvo['seq_num']:
                    continue

            score_afinidade, detalhes, sim = calcular_afinidade_tematica(alvo, cand)

            is_matriz_estruturante = any(r[0] == c_cod for r in matriz_pos.get(cod_alvo, []))
            if is_matriz_estruturante:
                score_afinidade += 2.0
                detalhes.insert(0, 'progressão cognitiva estruturante da BNCC')

            if cand['bloom_lvl'] >= bloom_alvo:
                score_afinidade += 0.15

            if score_afinidade >= 0.50 or is_matriz_estruturante:
                motivo = ''
                if is_matriz_estruturante:
                    motivo = f'Serve de degrau estruturante na progressão curricular: a consolidação de {cod_alvo} ({alvo["verbo"]}) é essencial para que o estudante possa avançar em {c_cod} ({cand["verbo"]}).'
                else:
                    dif_anos = cand['t_max'] - t_max
                    if dif_anos > 0:
                        motivo = f'Degrau formativo para o {cand["ano_fase"]}: a aprendizagem de {cod_alvo} consolida os fundamentos conceituais que serão expandidos em {c_cod} com foco em {cand["objetos_conhecimento"][0] if cand["objetos_conhecimento"] else "tema correlato"}.'
                    else:
                        motivo = f'Desdobramento no mesmo ciclo ({alvo["ano_fase"]}): esta habilidade fornece a bagagem inicial para o aprofundamento mais complexo exigido por {c_cod}.'

                candidatos_prepara.append({
                    'codigo': c_cod,
                    'score': score_afinidade,
                    'justificativa': motivo
                })

        candidatos_prepara.sort(key=lambda x: x['score'], reverse=True)
        final_prepara = [{'codigo': item['codigo'], 'justificativa': item['justificativa']} for item in candidatos_prepara[:3]]

        # --------------------------------------------------------------
        # C) RELACIONADAS (Interdisciplinares / Transversais - Máx 5)
        # --------------------------------------------------------------
        candidatos_relacionadas = []
        candidatos_etapa = por_etapa.get(alvo['etapa'], [])

        for outro_cod in candidatos_etapa:
            if outro_cod == cod_alvo:
                continue
            outro_h = habilidades[outro_cod]

            # Deve ser de OUTRO componente
            if outro_h['componente'] == comp:
                continue

            # Deve sobrepor o ano/fase escolar (mesmo período de desenvolvimento)
            sobreposicao_anos = (max(t_min, outro_h['t_min']) <= min(t_max, outro_h['t_max']))
            if not sobreposicao_anos:
                continue

            score_inter, detalhes, sim = calcular_afinidade_tematica(alvo, outro_h)

            # Requisito de interdisciplinaridade consistente
            # Similaridade vetorial >= 0.48 ou sobreposição de termos-chave conceituais
            if sim >= 0.48 or score_inter >= 0.52:
                motivo = f'Articulação interdisciplinar no {alvo["ano_fase"]} ({comp} e {outro_h["componente"]}): ambas conectam-se pelo tema de "{detalhes[0] if detalhes else "objetos correlatos"}", permitindo projetos integradores e aulas transversais.'
                candidatos_relacionadas.append({
                    'codigo': outro_cod,
                    'score': score_inter,
                    'justificativa': motivo
                })

        candidatos_relacionadas.sort(key=lambda x: x['score'], reverse=True)
        final_relacionadas = [{'codigo': item['codigo'], 'justificativa': item['justificativa']} for item in candidatos_relacionadas[:5]]

        total_requisitos += len(final_requisitos)
        total_prepara += len(final_prepara)
        total_relacionadas += len(final_relacionadas)

        resultado_final[cod_alvo] = {
            'descricao': alvo['descricao'],
            'componente': alvo['componente'],
            'ano_fase': alvo['ano_fase'],
            'requisitos': final_requisitos,
            'prepara_para': final_prepara,
            'relacionadas': final_relacionadas,
            'objetos_conhecimento': alvo['objetos_conhecimento']
        }

    # ==================================================================
    # 6. Validação e Assertions Anti-Alucinação
    # ==================================================================
    print('4. Executando validação anti-alucinação e integridade referencial...')
    assert len(resultado_final) == len(valid_codes), f'Esperado {len(valid_codes)}, gerado {len(resultado_final)}'

    erros_codigos = 0
    erros_temporais = 0
    erros_tamanho = 0

    for cod, dados in resultado_final.items():
        t_alvo = habilidades[cod]

        # Checa requisitos
        if len(dados['requisitos']) > 3:
            erros_tamanho += 1
        for r in dados['requisitos']:
            rcod = r['codigo']
            if rcod not in valid_codes or rcod == cod:
                erros_codigos += 1
            if habilidades[rcod]['t_min'] > t_alvo['t_min'] or habilidades[rcod]['t_max'] > t_alvo['t_max']:
                erros_temporais += 1

        # Checa prepara_para
        if len(dados['prepara_para']) > 3:
            erros_tamanho += 1
        for p in dados['prepara_para']:
            pcod = p['codigo']
            if pcod not in valid_codes or pcod == cod:
                erros_codigos += 1
            if habilidades[pcod]['t_max'] < t_alvo['t_max'] or habilidades[pcod]['t_min'] < t_alvo['t_min']:
                erros_temporais += 1

        # Checa relacionadas
        if len(dados['relacionadas']) > 5:
            erros_tamanho += 1
        for rel in dados['relacionadas']:
            relcod = rel['codigo']
            if relcod not in valid_codes or relcod == cod:
                erros_codigos += 1

    print(f'   -> Erros de códigos inexistentes ou auto-referência: {erros_codigos}')
    print(f'   -> Violações temporais estritas: {erros_temporais}')
    print(f'   -> Violações de limites máximos (3/3/5): {erros_tamanho}')

    assert erros_codigos == 0, 'Falha: códigos alucinados detectados!'
    assert erros_temporais == 0, 'Falha: violações temporais detectadas!'
    assert erros_tamanho == 0, 'Falha: limites excedidos!'

    # ==================================================================
    # 7. Gravação do Arquivo Final
    # ==================================================================
    print(f'5. Salvando arquivo final em: {OUTPUT_JSON}...')
    with open(OUTPUT_JSON, 'w', encoding='utf-8') as f:
        json.dump(resultado_final, f, ensure_ascii=False, indent=2)

    tamanho_bytes = os.path.getsize(OUTPUT_JSON)
    print(f'   Arquivo gerado com sucesso! Tamanho: {tamanho_bytes:,} bytes ({tamanho_bytes / (1024*1024):.2f} MB)')
    print('=' * 80)
    print('ESTATÍSTICAS CONSOLIDADAS:')
    print(f'- Total de Habilidades Mapeadas: {len(resultado_final):,}')
    print(f'- Total de Ligações de Requisitos: {total_requisitos:,} (média: {total_requisitos/len(resultado_final):.1f} por habilidade)')
    print(f'- Total de Ligações de Prepara Para: {total_prepara:,} (média: {total_prepara/len(resultado_final):.1f} por habilidade)')
    print(f'- Total de Ligações Interdisciplinares (Relacionadas): {total_relacionadas:,} (média: {total_relacionadas/len(resultado_final):.1f} por habilidade)')
    print('=' * 80)

if __name__ == '__main__':
    main()
