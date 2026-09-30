# -*- coding: utf-8 -*-
"""
integrar_referencias_reuna.py
Processa e integra os materiais de referência oficiais autorizados do Instituto Reúna:
1. Mapas de Foco BNCC (Ensino Fundamental - 1º ao 9º ano)
2. BNCC Comentada para o Ensino Médio (Linguagens, Matemática, Ciências da Natureza e Ciências Humanas)
3. Dimensões e Desenvolvimento das Competências Gerais da BNCC (Instituto Reúna / Instituto Ayrton Senna / Fundação Lemann)

Gera e atualiza:
- competencias_bncc.json (com dimensões, subdimensões e marcos de progressão nas 4 etapas)
- bncc_vetores.json (com comentários pedagógicos, objetos de conhecimento, contribuição dos componentes e classificação AF/AC)
- bncc_progressao_pre_requisitos.json (com todas as aprendizagens focais, complementares e comentários)
- relacoes_bncc_final.json (com objetos de conhecimento e dados enriquecidos para o EM)
"""

import openpyxl
import json
import os
import re
import glob

RECURSOS_DIR = 'recursos_adicionais'
COMPETENCIAS_XLSX = os.path.join(RECURSOS_DIR, 'BNCC_Competencias_Progressao.xlsx')
MAPAS_FOCO_XLSX = os.path.join(RECURSOS_DIR, 'MapasDeFocoBncc_Unificados.xlsx')
VETORES_JSON = 'bncc_vetores.json'
COMPETENCIAS_JSON = 'competencias_bncc.json'
PROGRESSAO_JSON = 'bncc_progressao_pre_requisitos.json'
RELACOES_JSON = 'relacoes_bncc_final.json'

def clean_text(val):
    if val is None:
        return ''
    t = str(val).strip()
    t = re.sub(r'\s+', ' ', t)
    return t

# ======================================================================
# 1. Extração das Dimensões e Marcos das Competências Gerais (Reúna / IAS)
# ======================================================================
def extrair_competencias_gerais_reuna():
    print(f'1. Extraindo matriz de progressão das Competências Gerais de {COMPETENCIAS_XLSX}...')
    if not os.path.exists(COMPETENCIAS_XLSX):
        print(f'   [AVISO] Arquivo {COMPETENCIAS_XLSX} não encontrado.')
        return {}

    wb = openpyxl.load_workbook(COMPETENCIAS_XLSX, data_only=True)
    dimensoes_por_comp = {}

    for sname in wb.sheetnames:
        s = wb[sname]
        c1 = clean_text(s.cell(1, 1).value)
        c2 = clean_text(s.cell(2, 1).value)
        c3 = clean_text(s.cell(3, 1).value)

        m_comp = re.search(r'Compet[êe]ncia:\s*(\d+)', c1, re.IGNORECASE)
        m_dim = re.search(r'Dimens[ãa]o:\s*(.+)', c2, re.IGNORECASE)
        m_sub = re.search(r'Subdimens[ãa]o:\s*(.+)', c3, re.IGNORECASE)

        if m_comp and m_dim and m_sub:
            comp_num = int(m_comp.group(1))
            dim_nome = clean_text(m_dim.group(1))
            sub_nome = clean_text(m_sub.group(1))

            marcos = {}
            for col in range(2, 7):
                header = clean_text(s.cell(4, col).value)
                desc = clean_text(s.cell(5, col).value)
                if header and desc and header.lower() != 'none':
                    # Normaliza a chave da etapa
                    etapa_key = 'ate_3_ef'
                    if '6' in header:
                        etapa_key = 'ate_6_ef'
                    elif '9' in header:
                        etapa_key = 'ate_9_ef'
                    elif 'médio' in header.lower() or 'medio' in header.lower() or 'mdio' in header.lower():
                        etapa_key = 'ate_3_em'

                    marcos[etapa_key] = {
                        'etapa_label': header,
                        'descricao_marco': desc
                    }

            if comp_num not in dimensoes_por_comp:
                dimensoes_por_comp[comp_num] = []

            dimensoes_por_comp[comp_num].append({
                'dimensao': dim_nome,
                'subdimensao': sub_nome,
                'marcos': marcos
            })

    print(f'   Sucesso! {len(dimensoes_por_comp)} Competências Gerais estruturadas com dimensões e marcos.')
    return dimensoes_por_comp

# ======================================================================
# 2. Extração da BNCC Comentada para o Ensino Médio (Instituto Reúna)
# ======================================================================
def extrair_ensino_medio_reuna():
    print('2. Extraindo BNCC Comentada para o Ensino Médio do Instituto Reúna...')
    em_files = glob.glob(os.path.join(RECURSOS_DIR, '*tecnologias.xlsx'))
    dados_em = {}

    for f in em_files:
        area_nome = os.path.basename(f).replace('.xlsx', '')
        wb = openpyxl.load_workbook(f, data_only=True)
        for sname in wb.sheetnames:
            s = wb[sname]
            # Cabeçalhos na linha 5, dados a partir da linha 6
            for r in range(6, s.max_row + 1):
                val_cod = clean_text(s.cell(r, 1).value)
                match = re.search(r'\b(EM13[A-Z0-9]+)\b', val_cod)
                if not match:
                    continue
                codigo = match.group(1).strip()
                comentario = clean_text(s.cell(r, 2).value)
                objeto = clean_text(s.cell(r, 3).value)
                componentes = clean_text(s.cell(r, 4).value)
                objetivos = clean_text(s.cell(r, 5).value)
                possibilidades = clean_text(s.cell(r, 6).value)

                dados_em[codigo] = {
                    'codigo': codigo,
                    'area_reuna': area_nome,
                    'competencia_especifica_reuna': sname,
                    'comentario': comentario,
                    'objeto_conhecimento': objeto,
                    'componentes_comentam': componentes,
                    'objetivos_aprendizagem': objetivos,
                    'possibilidades': possibilidades,
                    'referencia_fonte': 'Instituto Reúna - BNCC Comentada para o Ensino Médio'
                }

    print(f'   Sucesso! {len(dados_em)} habilidades do Ensino Médio extraídas com comentários e interdisciplinaridade.')
    return dados_em

# ======================================================================
# 3. Extração dos Mapas de Foco do Ensino Fundamental (Instituto Reúna)
# ======================================================================
def extrair_mapas_de_foco_ef():
    print(f'3. Extraindo classificações e objetivos dos Mapas de Foco de {MAPAS_FOCO_XLSX}...')
    if not os.path.exists(MAPAS_FOCO_XLSX):
        print(f'   [AVISO] Arquivo {MAPAS_FOCO_XLSX} não encontrado.')
        return {}

    wb = openpyxl.load_workbook(MAPAS_FOCO_XLSX, data_only=True)
    code_pattern = re.compile(r'\b(E[IF]\d{2}[A-Z]{2,3}\d{2})\b')
    dados_foco = {}

    for sname in wb.sheetnames:
        s = wb[sname]
        col_cod = None
        col_class = None
        col_obj = None
        col_prev = None
        col_coment = None
        header_row = None

        for r in range(1, min(7, s.max_row + 1)):
            for c in range(1, min(15, s.max_column + 1)):
                val = clean_text(s.cell(r, c).value).lower()
                if ('código' in val or 'cdigo' in val or 'habilidade' in val) and col_cod is None:
                    col_cod = c
                    header_row = r
                elif 'classifica' in val and col_class is None:
                    col_class = c
                elif 'objetivo' in val and col_obj is None:
                    col_obj = c
                elif 'conhecimento pr' in val and col_prev is None:
                    col_prev = c
                elif 'coment' in val and col_coment is None:
                    col_coment = c

            if col_cod and col_class:
                break

        if not (col_cod and header_row):
            continue

        for r in range(header_row + 1, s.max_row + 1):
            val_cod = clean_text(s.cell(r, col_cod).value)
            codes = code_pattern.findall(val_cod.upper())
            if not codes:
                continue

            for cod in codes:
                classif = clean_text(s.cell(r, col_class).value) if col_class else ''
                obj_apren = clean_text(s.cell(r, col_obj).value) if col_obj else ''
                coment = clean_text(s.cell(r, col_coment).value) if col_coment else ''

                # Normalização da classificação Reúna
                classif_norm = ''
                if classif.upper() in ['AF', 'FOCO', 'APRENDIZAGEM FOCAL']:
                    classif_norm = 'AF'
                elif classif.upper() in ['AC', 'COMPLEMENTAR', 'APRENDIZAGEM COMPLEMENTAR']:
                    classif_norm = 'AC'
                elif classif.upper() in ['APOIO', 'HABILIDADE DE APOIO']:
                    classif_norm = 'Apoio'

                if cod not in dados_foco:
                    dados_foco[cod] = {
                        'classificacao': classif_norm,
                        'objetivos_aprendizagem': obj_apren,
                        'comentario_reuna': coment,
                        'referencia_fonte': 'Instituto Reúna / Fundação Lemann / Itaú Social - Mapas de Foco BNCC'
                    }
                else:
                    if classif_norm and not dados_foco[cod]['classificacao']:
                        dados_foco[cod]['classificacao'] = classif_norm
                    if obj_apren and not dados_foco[cod]['objetivos_aprendizagem']:
                        dados_foco[cod]['objetivos_aprendizagem'] = obj_apren
                    if coment and not dados_foco[cod]['comentario_reuna']:
                        dados_foco[cod]['comentario_reuna'] = coment

    print(f'   Sucesso! {len(dados_foco)} habilidades do Ensino Fundamental com inteligência dos Mapas de Foco.')
    return dados_foco

# ======================================================================
# 4. Atualização de competencias_bncc.json
# ======================================================================
def atualizar_competencias_json(dimensoes_reuna):
    print(f'4. Atualizando {COMPETENCIAS_JSON} com dimensões, subdimensões e marcos de progressão...')
    if not os.path.exists(COMPETENCIAS_JSON):
        print(f'   [ERRO] {COMPETENCIAS_JSON} não encontrado.')
        return

    with open(COMPETENCIAS_JSON, 'r', encoding='utf-8') as f:
        data = json.load(f)

    # Atualiza as 10 Competências Gerais
    cg_dict = data.get('competencias_gerais', {})
    for key, cg in cg_dict.items():
        num = cg.get('numero')
        if num in dimensoes_reuna:
            cg['dimensoes_reuna'] = dimensoes_reuna[num]
            cg['referencia_pedagogica'] = 'Instituto Reúna, Instituto Ayrton Senna e Fundação Lemann (Dimensões e Desenvolvimento das Competências Gerais da BNCC)'

    # Metadados de Crédito e Atribuição Institucional
    data['creditos_institucionais'] = {
        'fonte_primaria': 'Base Nacional Comum Curricular (BNCC / MEC)',
        'referencia_progressao': 'Instituto Reúna (institutoreuna.org.br)',
        'obras_adaptadas': [
            'Dimensões e Desenvolvimento das Competências Gerais da BNCC (Instituto Reúna, Instituto Ayrton Senna e Fundação Lemann)',
            'Mapas de Foco da BNCC (Instituto Reúna, Fundação Lemann e Itaú Social)',
            'BNCC Comentada para o Ensino Médio (Instituto Reúna)'
        ],
        'nota_atribuicao': 'Material adaptado a partir das produções técnicas e metodologias de referência do Instituto Reúna sob autorização institucional para fins educacionais abertos.'
    }

    with open(COMPETENCIAS_JSON, 'w', encoding='utf-8') as f:
        json.dump(data, f, ensure_ascii=False, indent=2)

    print(f'   Sucesso! {COMPETENCIAS_JSON} salvo ({os.path.getsize(COMPETENCIAS_JSON)} bytes).')

# ======================================================================
# 5. Atualização de bncc_vetores.json
# ======================================================================
def atualizar_vetores_json(dados_em_reuna, dados_foco_ef):
    print(f'5. Atualizando {VETORES_JSON} com inteligência pedagógica do Instituto Reúna...')
    if not os.path.exists(VETORES_JSON):
        print(f'   [ERRO] {VETORES_JSON} não encontrado.')
        return

    with open(VETORES_JSON, 'r', encoding='utf-8') as f:
        vetores_data = json.load(f)

    cont_em = 0
    cont_ef = 0

    for item in vetores_data:
        cod = item.get('codigo', '').strip().upper()

        # Se for do Ensino Médio, aplica a BNCC Comentada Reúna
        if cod.startswith('EM') and cod in dados_em_reuna:
            em_info = dados_em_reuna[cod]
            if not item.get('comentario') and em_info['comentario']:
                item['comentario'] = em_info['comentario']
            if not item.get('objeto_conhecimento') and em_info['objeto_conhecimento']:
                item['objeto_conhecimento'] = em_info['objeto_conhecimento']
            if not item.get('possibilidades') and em_info['possibilidades']:
                item['possibilidades'] = em_info['possibilidades']
            
            item['componentes_comentam'] = em_info.get('componentes_comentam', '')
            item['objetivos_aprendizagem'] = em_info.get('objetivos_aprendizagem', '')
            item['fonte_pedagogica_reuna'] = em_info.get('referencia_fonte', '')
            cont_em += 1

        # Se for do Ensino Fundamental, aplica os Mapas de Foco Reúna
        if cod in dados_foco_ef:
            foco_info = dados_foco_ef[cod]
            if foco_info.get('classificacao'):
                item['classificacao_reuna'] = foco_info['classificacao']
            if foco_info.get('objetivos_aprendizagem') and not item.get('objetivos_aprendizagem'):
                item['objetivos_aprendizagem'] = foco_info['objetivos_aprendizagem']
            if foco_info.get('comentario_reuna') and (not item.get('comentario') or len(item.get('comentario', '')) < 10):
                item['comentario'] = foco_info['comentario_reuna']
            item['fonte_pedagogica_reuna'] = foco_info.get('referencia_fonte', '')
            cont_ef += 1

    with open(VETORES_JSON, 'w', encoding='utf-8') as f:
        json.dump(vetores_data, f, ensure_ascii=False, indent=2)

    print(f'   Sucesso! {cont_em} habilidades do EM e {cont_ef} do EF enriquecidas em {VETORES_JSON}.')

# ======================================================================
# 6. Atualização de bncc_progressao_pre_requisitos.json
# ======================================================================
def atualizar_progressao_json(dados_em_reuna, dados_foco_ef):
    print(f'6. Atualizando {PROGRESSAO_JSON}...')
    if not os.path.exists(PROGRESSAO_JSON):
        prog_data = {}
    else:
        with open(PROGRESSAO_JSON, 'r', encoding='utf-8') as f:
            prog_data = json.load(f)

    # Atualiza EF
    for cod, foco in dados_foco_ef.items():
        if cod not in prog_data:
            prog_data[cod] = {
                'classificacao': foco.get('classificacao', ''),
                'unidade': '',
                'pre_mesmo': [],
                'pre_ant': [],
                'apoio': [],
                'post_mesmo': [],
                'post_seg': [],
                'comentario': foco.get('comentario_reuna', ''),
                'objetivos_aprendizagem': foco.get('objetivos_aprendizagem', ''),
                'cadeias': []
            }
        else:
            if foco.get('classificacao') and not prog_data[cod].get('classificacao'):
                prog_data[cod]['classificacao'] = foco['classificacao']
            if foco.get('objetivos_aprendizagem'):
                prog_data[cod]['objetivos_aprendizagem'] = foco['objetivos_aprendizagem']
            if foco.get('comentario_reuna') and not prog_data[cod].get('comentario'):
                prog_data[cod]['comentario'] = foco['comentario_reuna']

    # Adiciona Ensino Médio na base de progressão
    for cod, em in dados_em_reuna.items():
        if cod not in prog_data:
            prog_data[cod] = {
                'classificacao': 'EM',
                'unidade': em.get('competencia_especifica_reuna', ''),
                'pre_mesmo': [],
                'pre_ant': [],
                'apoio': [],
                'post_mesmo': [],
                'post_seg': [],
                'comentario': em.get('comentario', ''),
                'componentes_comentam': em.get('componentes_comentam', ''),
                'objetivos_aprendizagem': em.get('objetivos_aprendizagem', ''),
                'possibilidades': em.get('possibilidades', ''),
                'cadeias': []
            }
        else:
            prog_data[cod]['comentario'] = em.get('comentario', '')
            prog_data[cod]['componentes_comentam'] = em.get('componentes_comentam', '')
            prog_data[cod]['objetivos_aprendizagem'] = em.get('objetivos_aprendizagem', '')
            prog_data[cod]['possibilidades'] = em.get('possibilidades', '')

    with open(PROGRESSAO_JSON, 'w', encoding='utf-8') as f:
        json.dump(prog_data, f, ensure_ascii=False, indent=2)

    print(f'   Sucesso! {len(prog_data)} habilidades com dados de progressão consolidados em {PROGRESSAO_JSON}.')

# ======================================================================
# 7. Atualização de relacoes_bncc_final.json
# ======================================================================
def atualizar_relacoes_json(dados_em_reuna):
    print(f'7. Atualizando {RELACOES_JSON}...')
    if not os.path.exists(RELACOES_JSON):
        print(f'   [AVISO] {RELACOES_JSON} não encontrado.')
        return

    with open(RELACOES_JSON, 'r', encoding='utf-8') as f:
        rel_data = json.load(f)

    cont_obj = 0
    for cod, item in rel_data.items():
        if cod in dados_em_reuna:
            em = dados_em_reuna[cod]
            if not item.get('objetos_conhecimento') and em.get('objeto_conhecimento'):
                objs = [x.strip() for x in em['objeto_conhecimento'].split('.') if x.strip()]
                item['objetos_conhecimento'] = objs
                cont_obj += 1
            if em.get('componentes_comentam'):
                item['componentes_comentam'] = em['componentes_comentam']
            if em.get('possibilidades'):
                item['possibilidades_curriculo'] = em['possibilidades']

    with open(RELACOES_JSON, 'w', encoding='utf-8') as f:
        json.dump(rel_data, f, ensure_ascii=False, indent=2)

    print(f'   Sucesso! {cont_obj} habilidades do EM receberam objetos de conhecimento em {RELACOES_JSON}.')

def main():
    print('======================================================================')
    print('  INTEGRAÇÃO OFICIAL DOS MATERIAIS DE REFERÊNCIA DO INSTITUTO REÚNA')
    print('======================================================================')
    dimensoes_reuna = extrair_competencias_gerais_reuna()
    dados_em_reuna = extrair_ensino_medio_reuna()
    dados_foco_ef = extrair_mapas_de_foco_ef()

    atualizar_competencias_json(dimensoes_reuna)
    atualizar_vetores_json(dados_em_reuna, dados_foco_ef)
    atualizar_progressao_json(dados_em_reuna, dados_foco_ef)
    atualizar_relacoes_json(dados_em_reuna)
    print('\n[CONCLUÍDO COM SUCESSO] Todos os artefatos de dados foram enriquecidos!')

if __name__ == '__main__':
    main()
